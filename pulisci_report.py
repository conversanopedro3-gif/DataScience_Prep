import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

print("--- 🔮 GENERAZIONE AUTOMATICA EXCEL NOTINO EXTRA-LARGO ---")

# 1. Inseriamo i dati reali direttamente nel codice per evitare errori di file mancanti
dati_perfetti = [
    {"Marca": "Hugo Boss", "Fragranza / Nome Profumo": "BOSS Bottled Eau de Toilette per uomo", "Prezzo": 37.80, "Unità Vendute": 591},
    {"Marca": "Davidoff", "Fragranza / Nome Profumo": "Cool Water Intense Eau de Parfum per uomo", "Prezzo": 23.90, "Unità Vendute": 563},
    {"Marca": "Lattafa", "Fragranza / Nome Profumo": "Asad Eau de Parfum per uomo 100 ml", "Prezzo": 24.60, "Unità Vendute": 554},
    {"Marca": "Jean Paul Gaultier", "Fragranza / Nome Profumo": "Le Beau Le Parfum Intense Eau de Parfum per uomo", "Prezzo": 78.20, "Unità Vendute": 553},
    {"Marca": "Calvin Klein", "Fragranza / Nome Profumo": "CK Be Eau de Toilette unisex", "Prezzo": 17.70, "Unità Vendute": 540},
    {"Marca": "Afnan", "Fragranza / Nome Profumo": "9 PM Rebel Eau de Parfum unisex 100 ml", "Prezzo": 30.90, "Unità Vendute": 524},
    {"Marca": "Yves Saint Laurent", "Fragranza / Nome Profumo": "Y Eau de Parfum ricaricabile per uomo", "Prezzo": 64.20, "Unità Vendute": 521},
    {"Marca": "DIOR", "Fragranza / Nome Profumo": "Sauvage Eau de Parfum ricaricabile per uomo", "Prezzo": 84.60, "Unità Vendute": 479},
    {"Marca": "Valentino", "Fragranza / Nome Profumo": "Born In Roma Intense Uomo Eau de Parfum per uomo", "Prezzo": 70.10, "Unità Vendute": 462},
    {"Marca": "Armaf", "Fragranza / Nome Profumo": "Club de Nuit Sillage Eau de Parfum per uomo", "Prezzo": 22.70, "Unità Vendute": 435},
    {"Marca": "TOM FORD", "Fragranza / Nome Profumo": "Noir Extreme spray corpo profumato per uomo 150 ml", "Prezzo": 43.20, "Unità Vendute": 430},
    {"Marca": "Lattafa", "Fragranza / Nome Profumo": "Khamrah Qahwa Eau de Parfum unisex 100 ml", "Prezzo": 26.00, "Unità Vendute": 376},
    {"Marca": "Rabanne", "Fragranza / Nome Profumo": "Invictus Eau de Toilette per uomo", "Prezzo": 53.70, "Unità Vendute": 364},
    {"Marca": "Dolce&Gabbana", "Fragranza / Nome Profumo": "Light Blue Pour Homme Eau de Toilette Eau de Toilette per uomo", "Prezzo": 65.80, "Unità Vendute": 353},
    {"Marca": "Armani", "Fragranza / Nome Profumo": "Emporio Stronger With You Intensely Eau de Parfum per uomo", "Prezzo": 56.30, "Unità Vendute": 325},
    {"Marca": "Lattafa", "Fragranza / Nome Profumo": "Eclaire Eau de Parfum unisex 100 ml", "Prezzo": 30.90, "Unità Vendute": 283},
    {"Marca": "Afnan", "Fragranza / Nome Profumo": "9 PM Night Out estratto profumato unisex 100 ml", "Prezzo": 49.90, "Unità Vendute": 273},
    {"Marca": "Calvin Klein", "Fragranza / Nome Profumo": "CK One Eau de Toilette unisex", "Prezzo": 29.20, "Unità Vendute": 254},
    {"Marca": "Montblanc", "Fragranza / Nome Profumo": "Explorer Eau de Parfum per uomo", "Prezzo": 30.80, "Unità Vendute": 253},
    {"Marca": "Afnan", "Fragranza / Nome Profumo": "Supremacy Collector's Edition Eau de Parfum per uomo 100 ml", "Prezzo": 45.30, "Unità Vendute": 239},
    {"Marca": "Armani", "Fragranza / Nome Profumo": "Acqua di Giò Profondo Eau de Parfum ricaricabile per uomo", "Prezzo": 64.40, "Unità Vendute": 209},
    {"Marca": "Lattafa", "Fragranza / Nome Profumo": "Khamrah Eau de Parfum unisex 100 ml", "Prezzo": 27.90, "Unità Vendute": 201},
    {"Marca": "Nasomatto", "Fragranza / Nome Profumo": "Black Afgano estratto profumato unisex 30 ml", "Prezzo": 110.60, "Unità Vendute": 167},
    {"Marca": "Montblanc", "Fragranza / Nome Profumo": "Explorer Extreme profumo per uomo 100 ml", "Prezzo": 45.70, "Unità Vendute": 163},
    {"Marca": "Afnan", "Fragranza / Nome Profumo": "Supremacy Not Only Intense estratto profumato per uomo 100 ml", "Prezzo": 37.80, "Unità Vendute": 140},
    {"Marca": "Armaf", "Fragranza / Nome Profumo": "Odyssey Mandarin Sky Eau de Parfum per uomo", "Prezzo": 23.40, "Unità Vendute": 121},
    {"Marca": "Jean Paul Gaultier", "Fragranza / Nome Profumo": "Scandal Pour Homme Le Parfum Eau de Parfum ricaricabile per uomo", "Prezzo": 83.20, "Unità Vendute": 117},
    {"Marca": "Lalique", "Fragranza / Nome Profumo": "Encre Noire Eau de Toilette per uomo", "Prezzo": 20.60, "Unità Vendute": 83},
    {"Marca": "Prada", "Fragranza / Nome Profumo": "Paradigme Eau de Parfum ricaricabile per uomo", "Prezzo": 66.90, "Unità Vendute": 78},
    {"Marca": "Armani", "Fragranza / Nome Profumo": "Acqua di Giò Eau de Toilette ricaricabile per uomo", "Prezzo": 56.50, "Unità Vendute": 34},
    {"Marca": "Dolce&Gabbana", "Fragranza / Nome Profumo": "The One For Men Eau de Toilette Eau de Toilette per uomo", "Prezzo": 59.20, "Unità Vendute": 17}
]

df_perfetto = pd.DataFrame(dati_perfetti)

# Generiamo anche il CSV locale così ripristiniamo il file mancante!
df_perfetto.to_csv("Report_Notino_PERFETTO.csv", index=False)

# 🚀 APPLICAZIONE STILE EXCEL EXTRA-LARGO E PROFESSIONALE
percorso_excel = "Report_Notino_PERFETTO.xlsx"

with pd.ExcelWriter(percorso_excel, engine='openpyxl') as writer:
    df_perfetto.to_excel(writer, sheet_name='Dati Puliti', index=False)
    
    worksheet = writer.sheets['Dati Puliti']
    worksheet.views.sheetView[0].showGridLines = True  # Griglia sempre visibile
    
    # Colori Tema Blu Notte Executive
    blu_header = "1B365D"
    bianco = "FFFFFF"
    grigio_chiaro = "D9D9D9"
    riga_zebra = "F9FBFC"
    
    font_header = Font(name="Segoe UI", size=11, bold=True, color=bianco)
    font_dati = Font(name="Segoe UI", size=10, color="000000")
    fill_header = PatternFill(start_color=blu_header, end_color=blu_header, fill_type="solid")
    fill_zebra = PatternFill(start_color=riga_zebra, end_color=riga_zebra, fill_type="solid")
    
    bordo_sottile = Border(
        left=Side(style='thin', color=grigio_chiaro), right=Side(style='thin', color=grigio_chiaro),
        top=Side(style='thin', color=grigio_chiaro), bottom=Side(style='thin', color=grigio_chiaro)
    )
    
    # 1. Formattazione Titoli Colonne
    worksheet.row_dimensions[1].height = 28
    for col_idx in range(1, 5):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = bordo_sottile
        
    # 2. Formattazione Righe Dati (Altezza e allineamenti)
    for r_idx in range(2, len(df_perfetto) + 2):
        worksheet.row_dimensions[r_idx].height = 22  # Diamo respiro verticale
        
        c_marca = worksheet.cell(row=r_idx, column=1)
        c_profumo = worksheet.cell(row=r_idx, column=2)
        c_prezzo = worksheet.cell(row=r_idx, column=3)
        c_vendite = worksheet.cell(row=r_idx, column=4)
        
        c_marca.alignment = Alignment(horizontal="left", vertical="center")
        c_profumo.alignment = Alignment(horizontal="left", vertical="center")
        c_prezzo.alignment = Alignment(horizontal="right", vertical="center")
        c_vendite.alignment = Alignment(horizontal="right", vertical="center")
        
        # Formati numerici reali
        c_prezzo.number_format = '#,##0.00" €"'
        c_vendite.number_format = '#,##0'
        
        for cell in [c_marca, c_profumo, c_prezzo, c_vendite]:
            cell.font = font_dati
            cell.border = bordo_sottile
            if r_idx % 2 == 1:
                cell.fill = fill_zebra

    # 3. IL TOCCO FINALE: Forziamo la colonna B a essere larghissima
    for col in worksheet.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter == 'B':
            worksheet.column_dimensions[col_letter].width = 65  # Colonna B spaziosa!
        else:
            max_len = max(len(str(cell.value or '')) for cell in col)
            worksheet.column_dimensions[col_letter].width = max(max_len + 5, 14)

print("🔥 SUCCESS_TOTAL! Il codice ha creato e organizzato l'Excel da zero!")
print("📦 File Excel generato con successo nella tua cartella: Report_Notino_PERFETTO.xlsx")